"""Generate a Word report with verbatim, version-pinned application source."""
from pathlib import Path
import hashlib
import json
import re
import subprocess
import zipfile
from lxml import etree
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DATA = json.loads((ROOT / "docs/manual-data.json").read_text(encoding="utf-8"))
SOURCE = DATA["source_commit"]
FILES = [
    "dist/index.html", "dist/styles.css", "dist/app.js", "dist/core.js",
    "dist/ocr.js", "dist/analyze.js", "dist/catalog.js", "build.cjs", "verify.cjs",
    "package.json", ".github/workflows/pages.yml", ".gitignore",
    "README.md", "package-lock.json",
]
HISTORICAL_COMMIT = "f3ec635a2a7c3d00588c1c139b1d1156721929a1"
HISTORICAL_FILE = ".github/workflows/import.yml"

def git_bytes(commit, path):
    return subprocess.check_output(["git", "show", commit + ":" + path], cwd=ROOT)

def add_link(paragraph, label, url):
    relation = paragraph.part.relate_to(
        url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), relation)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "14665C")
    props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    run.append(props)
    content = OxmlElement("w:t")
    content.text = label
    run.append(content)
    link.append(run)
    paragraph._p.append(link)

def rich(paragraph, value):
    offset = 0
    for match in re.finditer(r"https?://[^\s]+", value):
        paragraph.add_run(value[offset:match.start()])
        url = match.group(0).rstrip(".,;")
        add_link(paragraph, url, url)
        paragraph.add_run(match.group(0)[len(url):])
        offset = match.end()
    paragraph.add_run(value[offset:])
    return paragraph

def field(paragraph, name):
    run = paragraph.add_run()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), name)
    run._r.addnext(fld)

def table(doc, headers, rows):
    grid = doc.add_table(rows=1, cols=len(headers))
    grid.style = "Light Shading Accent 1"
    for cell, label in zip(grid.rows[0].cells, headers):
        cell.text = label
        for run in cell.paragraphs[0].runs:
            run.bold = True
    repeat = OxmlElement("w:tblHeader")
    grid.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for cell, value in zip(grid.add_row().cells, row):
            rich(cell.paragraphs[0], value)
    for row in grid.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(4)
                for run in paragraph.runs:
                    run.font.size = Pt(9)
    doc.add_paragraph()

doc = Document()
section = doc.sections[0]
section.page_height = Cm(29.7)
section.page_width = Cm(21)
section.top_margin = Cm(2)
section.bottom_margin = Cm(2)
section.left_margin = Cm(2)
section.right_margin = Cm(2)
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.12
for name in ["Title", "Heading 1", "Heading 2"]:
    doc.styles[name].font.name = "Calibri"
    doc.styles[name].font.color.rgb = RGBColor.from_string("123E38")
code = doc.styles.add_style("Source Code", 1)
code.font.name = "Consolas"
code.font.size = Pt(8)
code.paragraph_format.space_after = Pt(0)
code.paragraph_format.space_before = Pt(0)
code.paragraph_format.line_spacing = Pt(9)
code.paragraph_format.keep_with_next = False
code.paragraph_format.widow_control = False
section.header.paragraphs[0].text = "FONCEP | Informe técnico y guía de uso"
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run("Página ")
field(footer, "PAGE")
footer.add_run(" de ")
field(footer, "NUMPAGES")

doc.add_paragraph(DATA["title"], "Title")
doc.add_paragraph(DATA["subtitle"], "Subtitle")
doc.add_paragraph(DATA["date"])
doc.add_paragraph("Versión documentada: " + SOURCE)
rich(doc.add_paragraph(), "Aplicación: https://deivy199130-arch.github.io/foncep-organizador/")
doc.add_paragraph(
    "Este informe explica cómo se construyó y cómo se utiliza la aplicación; "
    "distingue el tratamiento local de los expedientes del alojamiento público del código."
)
doc.add_paragraph(
    "Incluye una sección independiente con todas las líneas de los archivos propios "
    "de la versión publicada y un anexo de dependencias y migración histórica."
)
doc.add_page_break()
doc.add_heading("Contenido", level=1)
for chapter in DATA["chapters"]:
    doc.add_paragraph(chapter["title"])
doc.add_paragraph("16. Paso independiente: código completo de la aplicación")
doc.add_paragraph("17. Anexo: dependencias exactas y migración histórica")
doc.add_paragraph("18. Integridad y alcance de los anexos")
doc.add_page_break()

for chapter in DATA["chapters"]:
    doc.add_heading(chapter["title"], level=1)
    for block in chapter["blocks"]:
        kind = block["kind"]
        if kind == "paragraph":
            rich(doc.add_paragraph(), block["text"])
        elif kind == "bullets":
            for item in block["items"]:
                rich(doc.add_paragraph(style="List Bullet"), item)
        elif kind == "steps":
            for index, item in enumerate(block["items"], 1):
                rich(doc.add_paragraph(), str(index) + ". " + item)
        elif kind == "table":
            table(doc, block["headers"], block["rows"])
        else:
            raise ValueError("Unsupported document block: " + kind)

appendix = doc.add_section(WD_SECTION.NEW_PAGE)
appendix.orientation = WD_ORIENT.LANDSCAPE
appendix.page_width = Cm(29.7)
appendix.page_height = Cm(21)
appendix.top_margin = Cm(1.5)
appendix.bottom_margin = Cm(1.5)
appendix.left_margin = Cm(1.5)
appendix.right_margin = Cm(1.5)
doc.add_heading("16. Paso independiente: código completo de la aplicación", level=1)
doc.add_paragraph(
    "Cada bloque reproduce íntegramente un archivo de la versión " + SOURCE + ". "
    "La numeración de seis dígitos y el separador | son ayudas de lectura añadidas "
    "para este informe y no pertenecen al código. Los saltos visuales por ancho de "
    "página no representan nuevas líneas del archivo. No se omiten cuerpos de funciones."
)
expected = []
manifest = []

def add_source(path, commit, title=None):
    raw = git_bytes(commit, path)
    source = raw.decode("utf-8")
    digest = hashlib.sha256(raw).hexdigest()
    doc.add_heading(title or path, level=2)
    doc.add_paragraph("SHA-256 del archivo: " + digest)
    rich(doc.add_paragraph(),
         "Fuente: https://github.com/deivy199130-arch/foncep-organizador/blob/"
         + commit + "/" + path)
    lines = source.split("\n")
    doc.add_paragraph(
        "Líneas del archivo: " + str(len(source.splitlines())) +
        ". El posible párrafo vacío final preserva un salto de línea de cierre."
    )
    for number, line in enumerate(lines, 1):
        value = f"{number:06d} | " + line
        doc.add_paragraph(value, style="Source Code")
        expected.append(value)
    manifest.append({
        "path": path, "commit": commit, "sha256": digest,
        "bytes": len(raw), "lines": len(source.splitlines()),
    })
    doc.add_page_break()

for path in FILES:
    if path == "package-lock.json":
        continue
    add_source(path, SOURCE)

doc.add_heading("17. Anexo: dependencias exactas y migración histórica", level=1)
doc.add_paragraph(
    "package-lock.json fue generado por npm y se incluye completo para conocer "
    "las versiones e integridades de los paquetes. No es código escrito manualmente. "
    "El flujo import.yml se recupera del historial y ya fue eliminado de main. "
    "Su URL antigua corresponde únicamente a la transferencia de una sola vez."
)
add_source("package-lock.json", SOURCE, "17.1. package-lock.json completo")
add_source(HISTORICAL_FILE, HISTORICAL_COMMIT,
           "17.2. import.yml histórico, eliminado de la versión vigente")

doc.add_heading("18. Integridad y alcance de los anexos", level=1)
doc.add_paragraph(
    "Se comprobará que el archivo Word puede abrirse como paquete DOCX válido "
    "y que cada párrafo de código coincide exactamente con la línea original "
    "obtenida de Git. Las huellas permiten cotejar las fuentes. Esta comprobación "
    "no prueba que el software esté libre de errores ni certifica sus dependencias."
)
table(doc, ["Archivo", "Versión", "Líneas / SHA-256"], [
    [entry["path"], entry["commit"],
     str(entry["lines"]) + " / " + entry["sha256"]]
    for entry in manifest
])
doc.add_paragraph(
    "Los archivos generados vendor no se atribuyen al desarrollo propio: "
    "su procedencia se explica en los capítulos 4 y 5 y en package-lock.json. "
    "El generador de este documento está disponible en scripts/generate_manual.py; "
    "es una herramienta de documentación y no forma parte del código ejecutado por la página."
)
doc.core_properties.title = DATA["title"] + " — Informe técnico y código"
doc.core_properties.subject = "Implementación, uso, GitHub, OCR y tratamiento de datos"
doc.core_properties.author = "Proyecto foncep-organizador"
doc.core_properties.keywords = "FONCEP, OCR, GitHub Pages, privacidad, código fuente"
destination = ROOT / "docs/FONCEP_Manual_tecnico_y_codigo.docx"
destination.parent.mkdir(exist_ok=True)
doc.save(destination)

with zipfile.ZipFile(destination) as archive:
    assert archive.testzip() is None, "Corrupt DOCX archive"
    for name in archive.namelist():
        if name.endswith(".xml"):
            etree.fromstring(archive.read(name))
reopened = Document(destination)
actual = [paragraph.text for paragraph in reopened.paragraphs
          if paragraph.style.name == "Source Code"]
assert actual == expected, "Source code appendix was truncated or altered"
assert len(reopened.tables) >= 7, "Missing explanatory tables"
assert len(DATA["chapters"]) == 15
assert destination.stat().st_size > 20000, "Document unexpectedly small"
(ROOT / "docs/manual-integrity.json").write_text(json.dumps({
    "application_commit": SOURCE,
    "source_files": manifest,
    "code_paragraphs": len(expected),
    "docx_sha256": hashlib.sha256(destination.read_bytes()).hexdigest(),
    "validation": "DOCX/XML valid; all code paragraphs match Git verbatim",
}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print("Word generado:", destination.name)
print("Bytes:", destination.stat().st_size)
print("Archivos de código:", len(manifest))
print("Párrafos de código completos y verificados:", len(expected))
